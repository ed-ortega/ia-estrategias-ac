import json
import pandas as pd

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


ENTRADA = "resultado_hvac.xlsx"
SALIDA = "comparacion_clima.xlsx"


# =========================================================
# LEER RESULTADOS
# =========================================================
df = pd.read_excel(ENTRADA)

registros = []

for _, row in df.iterrows():

    try:
        resultado = json.loads(row["resultado"])
    except Exception:
        continue

    sp = resultado.get("SP", {})
    spd1 = resultado.get("SPD1", {})
    spd2 = resultado.get("SPD2", {})

    origen = resultado.get(
        "origen_clima",
        "desconocido"
    )

    nombres = {
        "openmeteo": "Open-Meteo",
        "weatherapi": "WeatherAPI",
        "historico_hvac": "Parquet"
    }

    registros.append({
        "Sucursal": row.get("Sucursal"),
        "Ubicacion": row.get("Ubicacion"),
        "Tecnologia": row.get("Tecnologia"),
        "Estado": row.get("Estado"),
        "Ciudad": row.get("Ciudad"),
        "Latitud": row.get("Latitud"),
        "Longitud": row.get("Longitud"),

        "Fuente": nombres.get(origen, origen),

        "Clima SP": sp.get("clima"),
        "Clima SPD1": spd1.get("clima"),
        "Clima SPD2": spd2.get("clima"),

        "SP": sp.get("temp_prom"),
        "SPD1": spd1.get("temp_prom"),
        "SPD2": spd2.get("temp_prom"),
    })


detalle = pd.DataFrame(registros)


# =========================================================
# TEMPERATURAS NUMÉRICAS
# =========================================================
for columna in [
    "SP",
    "SPD1",
    "SPD2"
]:
    detalle[columna] = pd.to_numeric(
        detalle[columna],
        errors="coerce"
    )


# =========================================================
# RESUMEN POR FUENTE
# =========================================================
resumen = (
    detalle
    .groupby("Fuente")
    .agg(
        Registros=("Fuente", "size"),
        SP=("SP", "mean"),
        SPD1=("SPD1", "mean"),
        SPD2=("SPD2", "mean")
    )
    .reset_index()
)

for columna in ["SP", "SPD1", "SPD2"]:
    resumen[columna] = resumen[columna].round(2)


# =========================================================
# PROMEDIO OPEN-METEO
# =========================================================
openmeteo = detalle[
    detalle["Fuente"] == "Open-Meteo"
]

open_sp = round(
    openmeteo["SP"].mean(),
    2
)

open_spd1 = round(
    openmeteo["SPD1"].mean(),
    2
)

open_spd2 = round(
    openmeteo["SPD2"].mean(),
    2
)


# =========================================================
# CASOS WEATHERAPI / PARQUET
# =========================================================
especiales = detalle[
    detalle["Fuente"].isin([
        "WeatherAPI",
        "Parquet"
    ])
].copy()


# Etiquetas cortas
etiquetas = []
parquet_num = 1

for fuente in especiales["Fuente"]:

    if fuente == "WeatherAPI":
        etiquetas.append("WeatherAPI")

    else:
        etiquetas.append(
            f"Parquet {parquet_num}"
        )
        parquet_num += 1


especiales.insert(
    0,
    "Caso",
    etiquetas
)


# =========================================================
# TABLA VISUAL
# =========================================================
tabla_visual = especiales[[
    "Caso",
    "Fuente",
    "Sucursal",
    "Ubicacion",
    "Tecnologia",
    "Estado",
    "Ciudad",
    "Clima SP",
    "Clima SPD1",
    "Clima SPD2",
]].copy()


# =========================================================
# TABLA DE TEMPERATURAS
# =========================================================
tabla_temp = especiales[[
    "Caso",
    "SP",
    "SPD1",
    "SPD2"
]].copy()


tabla_temp["Dif SP"] = (
    tabla_temp["SP"] - open_sp
).round(2)

tabla_temp["Dif SPD1"] = (
    tabla_temp["SPD1"] - open_spd1
).round(2)

tabla_temp["Dif SPD2"] = (
    tabla_temp["SPD2"] - open_spd2
).round(2)


# =========================================================
# DATOS PARA GRÁFICA PRINCIPAL
# =========================================================
grafico = pd.DataFrame([
    {
        "Fuente": "Open-Meteo",
        "SP": open_sp,
        "SPD1": open_spd1,
        "SPD2": open_spd2
    }
])


for _, fila in especiales.iterrows():

    grafico.loc[len(grafico)] = {
        "Fuente": fila["Caso"],
        "SP": fila["SP"],
        "SPD1": fila["SPD1"],
        "SPD2": fila["SPD2"]
    }


# =========================================================
# DATOS PARA GRÁFICA DE DIFERENCIAS
# =========================================================
grafico_dif = tabla_temp[[
    "Caso",
    "Dif SP",
    "Dif SPD1",
    "Dif SPD2"
]].copy()


# =========================================================
# EXPORTAR
# =========================================================
with pd.ExcelWriter(
    SALIDA,
    engine="openpyxl"
) as writer:

    detalle.to_excel(
        writer,
        sheet_name="Detalle",
        index=False
    )

    resumen.to_excel(
        writer,
        sheet_name="Resumen",
        index=False
    )

    tabla_visual.to_excel(
        writer,
        sheet_name="Comparacion visual",
        index=False,
        startrow=4
    )

    tabla_temp.to_excel(
        writer,
        sheet_name="Comparacion visual",
        index=False,
        startrow=11
    )

    grafico.to_excel(
        writer,
        sheet_name="Comparacion visual",
        index=False,
        startrow=19
    )

    grafico_dif.to_excel(
        writer,
        sheet_name="Comparacion visual",
        index=False,
        startrow=27
    )


# =========================================================
# FORMATO
# =========================================================
wb = load_workbook(SALIDA)

AZUL = "1F4E78"
AZUL_CLARO = "D9EAF7"
GRIS = "F3F6F8"
VERDE = "E2F0D9"

BORDE = Side(
    style="thin",
    color="D9E1F2"
)


def estilo_encabezado(
    ws,
    fila,
    inicio,
    fin
):

    for col in range(
        inicio,
        fin + 1
    ):

        celda = ws.cell(
            row=fila,
            column=col
        )

        celda.font = Font(
            bold=True,
            color="FFFFFF"
        )

        celda.fill = PatternFill(
            "solid",
            fgColor=AZUL
        )

        celda.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


# =========================================================
# HOJA DETALLE
# =========================================================
ws_detalle = wb["Detalle"]

estilo_encabezado(
    ws_detalle,
    1,
    1,
    ws_detalle.max_column
)

ws_detalle.freeze_panes = "A2"
ws_detalle.auto_filter.ref = ws_detalle.dimensions


# =========================================================
# HOJA RESUMEN
# =========================================================
ws_resumen = wb["Resumen"]

estilo_encabezado(
    ws_resumen,
    1,
    1,
    ws_resumen.max_column
)


# ---------------------------------------------------------
# GRÁFICA TEMPERATURAS POR FUENTE
# ---------------------------------------------------------
chart_resumen = BarChart()

chart_resumen.type = "col"

chart_resumen.title = (
    "Temperatura promedio por fuente"
)

chart_resumen.y_axis.title = (
    "Temperatura (°F)"
)

chart_resumen.x_axis.title = (
    "Fuente"
)

datos = Reference(
    ws_resumen,
    min_col=3,
    max_col=5,
    min_row=1,
    max_row=ws_resumen.max_row
)

categorias = Reference(
    ws_resumen,
    min_col=1,
    min_row=2,
    max_row=ws_resumen.max_row
)

chart_resumen.add_data(
    datos,
    titles_from_data=True
)

chart_resumen.set_categories(
    categorias
)

chart_resumen.height = 9
chart_resumen.width = 18
chart_resumen.legend.position = "b"

ws_resumen.add_chart(
    chart_resumen,
    "G2"
)


# ---------------------------------------------------------
# GRÁFICA CANTIDAD DE REGISTROS
# ---------------------------------------------------------
chart_cantidad = BarChart()

chart_cantidad.type = "bar"

chart_cantidad.title = (
    "Registros procesados por fuente"
)

chart_cantidad.x_axis.title = (
    "Cantidad de registros"
)

chart_cantidad.y_axis.title = (
    "Fuente"
)

datos_cantidad = Reference(
    ws_resumen,
    min_col=2,
    min_row=1,
    max_row=ws_resumen.max_row
)

categorias_cantidad = Reference(
    ws_resumen,
    min_col=1,
    min_row=2,
    max_row=ws_resumen.max_row
)

chart_cantidad.add_data(
    datos_cantidad,
    titles_from_data=True
)

chart_cantidad.set_categories(
    categorias_cantidad
)

chart_cantidad.height = 7
chart_cantidad.width = 18

chart_cantidad.legend = None

ws_resumen.add_chart(
    chart_cantidad,
    "G20"
)


# =========================================================
# COMPARACION VISUAL
# =========================================================
ws = wb["Comparacion visual"]


# ---------------------------------------------------------
# TÍTULO
# ---------------------------------------------------------
ws.merge_cells("A1:J1")

ws["A1"] = (
    "Comparación de fuentes climáticas"
)

ws["A1"].font = Font(
    bold=True,
    size=18,
    color="FFFFFF"
)

ws["A1"].fill = PatternFill(
    "solid",
    fgColor=AZUL
)

ws["A1"].alignment = Alignment(
    horizontal="center",
    vertical="center"
)

ws.row_dimensions[1].height = 30


# ---------------------------------------------------------
# SUBTÍTULO
# ---------------------------------------------------------
ws.merge_cells("A2:J2")

ws["A2"] = (
    "Open-Meteo como referencia frente a "
    "WeatherAPI y predicciones del histórico HVAC"
)

ws["A2"].font = Font(
    italic=True,
    size=11
)

ws["A2"].alignment = Alignment(
    horizontal="center"
)


# =========================================================
# CASOS ANALIZADOS
# =========================================================
ws["A4"] = "Casos analizados"

ws["A4"].font = Font(
    bold=True,
    size=13,
    color=AZUL
)

estilo_encabezado(
    ws,
    5,
    1,
    10
)


# =========================================================
# TEMPERATURAS
# =========================================================
ws["A11"] = (
    "Temperaturas y diferencia contra Open-Meteo"
)

ws["A11"].font = Font(
    bold=True,
    size=13,
    color=AZUL
)

estilo_encabezado(
    ws,
    12,
    1,
    7
)


# =========================================================
# REFERENCIA OPEN-METEO
# =========================================================
ws["I11"] = "Referencia Open-Meteo"

ws["I11"].font = Font(
    bold=True,
    size=13,
    color=AZUL
)


referencia = [
    ["Periodo", "Promedio °F"],
    ["SP", open_sp],
    ["SPD1", open_spd1],
    ["SPD2", open_spd2],
]


for fila_idx, fila in enumerate(
    referencia,
    start=12
):

    for col_idx, valor in enumerate(
        fila,
        start=9
    ):

        ws.cell(
            row=fila_idx,
            column=col_idx,
            value=valor
        )


estilo_encabezado(
    ws,
    12,
    9,
    10
)


for fila in range(
    13,
    16
):

    for col in range(
        9,
        11
    ):

        ws.cell(
            row=fila,
            column=col
        ).fill = PatternFill(
            "solid",
            fgColor=VERDE
        )


# =========================================================
# DATOS GRÁFICA PRINCIPAL
# =========================================================
fila_grafico = 20

estilo_encabezado(
    ws,
    fila_grafico,
    1,
    4
)


# =========================================================
# GRÁFICA PRINCIPAL
# =========================================================
chart = BarChart()

chart.type = "col"

chart.title = (
    "Open-Meteo vs WeatherAPI vs Parquet"
)

chart.y_axis.title = (
    "Temperatura (°F)"
)

chart.x_axis.title = (
    "Fuente"
)

datos = Reference(
    ws,
    min_col=2,
    max_col=4,
    min_row=fila_grafico,
    max_row=fila_grafico + len(grafico)
)

categorias = Reference(
    ws,
    min_col=1,
    min_row=fila_grafico + 1,
    max_row=fila_grafico + len(grafico)
)

chart.add_data(
    datos,
    titles_from_data=True
)

chart.set_categories(
    categorias
)

chart.height = 12
chart.width = 24

chart.legend.position = "b"

ws.add_chart(
    chart,
    "F19"
)


# =========================================================
# SECCIÓN DIFERENCIAS
# =========================================================
ws["A27"] = (
    "Diferencia respecto a Open-Meteo"
)

ws["A27"].font = Font(
    bold=True,
    size=13,
    color=AZUL
)

ws["A28"] = (
    "Positivo = por encima | "
    "Negativo = por debajo"
)

ws["A28"].font = Font(
    italic=True,
    size=10
)


fila_dif = 29

estilo_encabezado(
    ws,
    fila_dif,
    1,
    4
)


# =========================================================
# GRÁFICA DIFERENCIAS
# =========================================================
if len(grafico_dif) > 0:

    chart_dif = BarChart()

    chart_dif.type = "col"

    chart_dif.title = (
        "¿Cuánto se aleja de Open-Meteo?"
    )

    chart_dif.y_axis.title = (
        "Diferencia (°F)"
    )

    chart_dif.x_axis.title = (
        "Caso"
    )

    datos_dif = Reference(
        ws,
        min_col=2,
        max_col=4,
        min_row=fila_dif,
        max_row=fila_dif + len(grafico_dif)
    )

    categorias_dif = Reference(
        ws,
        min_col=1,
        min_row=fila_dif + 1,
        max_row=fila_dif + len(grafico_dif)
    )

    chart_dif.add_data(
        datos_dif,
        titles_from_data=True
    )

    chart_dif.set_categories(
        categorias_dif
    )

    chart_dif.height = 10
    chart_dif.width = 20

    chart_dif.legend.position = "b"

    ws.add_chart(
        chart_dif,
        "F34"
    )


# =========================================================
# FORMATO DE FILAS
# =========================================================
for fila in range(
    6,
    6 + len(tabla_visual)
):

    for col in range(
        1,
        11
    ):

        celda = ws.cell(
            row=fila,
            column=col
        )

        celda.fill = PatternFill(
            "solid",
            fgColor=GRIS
        )

        celda.border = Border(
            bottom=BORDE
        )

        celda.alignment = Alignment(
            vertical="center"
        )


for fila in range(
    13,
    13 + len(tabla_temp)
):

    for col in range(
        1,
        8
    ):

        celda = ws.cell(
            row=fila,
            column=col
        )

        celda.border = Border(
            bottom=BORDE
        )


# =========================================================
# ANCHOS
# =========================================================
anchos = {
    "A": 14,
    "B": 14,
    "C": 22,
    "D": 24,
    "E": 14,
    "F": 17,
    "G": 17,
    "H": 15,
    "I": 15,
    "J": 15,
}

for col, ancho in anchos.items():

    ws.column_dimensions[
        col
    ].width = ancho


ws.freeze_panes = "A5"


# =========================================================
# GUARDAR
# =========================================================
wb.save(SALIDA)


print("\n✅ Reporte generado:")
print(SALIDA)

print("\n📊 Registros por fuente:")

print(
    detalle["Fuente"]
    .value_counts()
    .to_string()
)

print(
    "\n📈 Gráficas generadas:"
)

print(
    "1. Temperatura promedio por fuente"
)

print(
    "2. Cantidad de registros por fuente"
)

print(
    "3. Comparación Open-Meteo / WeatherAPI / Parquet"
)

print(
    "4. Diferencia respecto a Open-Meteo"
)